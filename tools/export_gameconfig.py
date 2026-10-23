#!/usr/bin/env python3
"""Export config/gameconfig/*.xlsx to client JSON.

Source of truth is the in-repo directory config/gameconfig (override with
--src or GAMECONFIG_SRC). Runtime client code only reads the JSON written to
client/assets/resources/config. xlsx is not loaded by Creator.

Header rows match the old ExcelReader layout:
  row 1  Chinese note
  row 2  export flag: 1 server, 2 client, 3 both, blank/0 skip
  row 3  field name
  row 4+ data

Columns without a 1/2/3 flag are designer comments and are not exported.
system_gate.json is derived from system.xlsx integer *_open_day / *_days keys.
opencondition.json keeps the client fields id/name/condition_type/level/
open_server_day and also keeps the raw columns. main_cafe is an alias of
main_pub_task, mail is an alias of main_mail. The nav id bag has no atom, so
a level-1 row is added for that id. Existing gate ids still resolve.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import sys
import uuid
import zipfile
from datetime import date, datetime, time
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write("missing openpyxl. pip install -r tools/requirements.txt\n")
    sys.exit(2)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRC = ROOT / "config" / "gameconfig"
DEFAULT_OUT = ROOT / "client" / "assets" / "resources" / "config"
FIELD_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
INT_RE = re.compile(r"^-?\d+$")

META_TEMPLATE = """{{
  "ver": "2.0.1",
  "importer": "json",
  "imported": true,
  "uuid": "{uuid}",
  "files": [
    ".json"
  ],
  "subMetas": {{}},
  "userData": {{}}
}}
"""


def main() -> int:
    parser = argparse.ArgumentParser(description="Export gameconfig xlsx to client JSON")
    parser.add_argument("--src", default=os.environ.get("GAMECONFIG_SRC", str(DEFAULT_SRC)))
    parser.add_argument("--out", default=os.environ.get("GAMECONFIG_OUT", str(DEFAULT_OUT)))
    args = parser.parse_args()
    src = Path(args.src)
    out = Path(args.out)
    if not src.is_dir():
        sys.stderr.write("gameconfig directory not found: %s\n" % src)
        return 2
    books = sorted(p for p in src.glob("*.xlsx") if not p.name.startswith("~$"))
    if not books:
        sys.stderr.write("no xlsx in %s\n" % src)
        return 2
    out.mkdir(parents=True, exist_ok=True)

    written: list[str] = []
    failed: list[str] = []
    extra = 0
    system_rows: list[dict] | None = None
    open_rows: list[dict] | None = None
    for index, path in enumerate(books, 1):
        if index == 1 or index % 40 == 0 or index == len(books):
            print("read %d/%d %s" % (index, len(books), path.name), flush=True)
        try:
            sheets = read_book(path)
        except Exception as err:
            failed.append("%s: %s" % (path.name, err))
            continue
        if not sheets:
            failed.append("%s: no header" % path.name)
            continue
        primary = pick_primary(path.stem, sheets)
        note = "exported from config/gameconfig/%s sheet %s. flag 1/2/3 columns only." % (
            path.name,
            primary["sheet"],
        )
        write_table(out, path.stem, note, primary["rows"])
        written.append(path.stem)
        if path.stem == "system":
            system_rows = loose_rows(path, primary["sheet"])
        if path.stem == "opencondition":
            open_rows = primary["rows"]
        for index, sheet in enumerate(sheets):
            if sheet is primary:
                continue
            if not sheet["rows"]:
                continue
            slug = sheet_slug(sheet["sheet"], index)
            name = "%s__%s" % (path.stem, slug)
            if name in written:
                name = "%s_%d" % (name, index)
            write_table(
                out,
                name,
                "exported from config/gameconfig/%s extra sheet %s." % (path.name, sheet["sheet"]),
                sheet["rows"],
            )
            written.append(name)
            extra += 1

    gates = build_system_gate(system_rows or [])
    write_table(
        out,
        "system_gate",
        "derived from config/gameconfig/system.xlsx. integer keys containing open_day or ending with _days. open_server_day is that value. remark is note.",
        gates,
    )
    if "system_gate" not in written:
        written.append("system_gate")

    projected = project_opencondition(open_rows or [], gates)
    write_table(
        out,
        "opencondition",
        "exported from config/gameconfig/opencondition.xlsx. id=atom, name=note, level=condition_value, open_server_day from system *_open_day when the key is {atom}_open_day. main_cafe aliases main_pub_task. mail aliases main_mail. bag is a level-1 nav row because the sheet has no bag atom. raw columns kept.",
        projected,
    )

    names = sorted(set(written))
    manifest = {
        "table": "manifest",
        "note": "client load list. source config/gameconfig. xlsx=%d json=%d extra_sheets=%d failed=%d%s"
        % (
            len(books),
            len(names),
            extra,
            len(failed),
            ("; " + "; ".join(failed)) if failed else "",
        ),
        "tables": names,
        "rows": [],
    }
    write_json(out / "manifest.json", manifest)
    ensure_meta(out / "manifest.json.meta")

    print(
        "xlsx=%d tables=%d extra_sheets=%d opencondition=%d system_gate=%d failed=%d"
        % (len(books), len(names), extra, len(projected), len(gates), len(failed))
    )
    for line in failed:
        print("FAIL", line)
    return 0 if not failed else 1


def read_book(path: Path) -> list[dict]:
    wb = open_book(path)
    found = []
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = [tuple(row) for row in ws.iter_rows(values_only=True)]
            parsed = parse_sheet(rows)
            if parsed is None:
                continue
            found.append({"sheet": sheet_name, "rows": parsed})
    finally:
        wb.close()
    return found


def open_book(path: Path):
    # Not read_only: several workbooks lie about their dimension, and read_only
    # then stops after the first cell.
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except Exception:
        # Some workbooks store zip entry names with backslashes.
        raw = path.read_bytes()
        buf = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as zin, zipfile.ZipFile(buf, "w") as zout:
            for info in zin.infolist():
                zout.writestr(info.filename.replace("\\", "/"), zin.read(info.filename))
        buf.seek(0)
        return openpyxl.load_workbook(buf, data_only=True)


def loose_rows(path: Path, sheet_name: str) -> list[dict]:
    wb = open_book(path)
    try:
        ws = wb[sheet_name]
        rows = [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return parse_sheet(rows, flagged_only=False) or []


TYPE_WORDS = {"string", "int", "integer", "json", "float", "number", "bool", "double", "long", "text"}


def parse_sheet(rows: list[tuple], flagged_only: bool = True) -> list[dict] | None:
    if len(rows) < 3:
        return None
    flags = [flag_of(cell) for cell in rows[1]]
    fields = list(rows[2])
    known = sum(1 for flag in flags if flag in (0, 1, 2, 3))
    other = sum(1 for flag in flags if flag == "other")
    type_words = sum(1 for cell in rows[1] if is_type_word(cell))
    idents = [(index, str(field).strip()) for index, field in enumerate(fields) if is_field(field)]
    type_header = type_words >= 1 and known == 0 and len(idents) >= 1
    if len(idents) < 1:
        return None
    if not type_header and (known < 1 or other > known):
        return None
    columns = []
    used: dict[str, int] = {}
    for index, name in idents:
        flag = flags[index] if index < len(flags) else None
        if flagged_only and not type_header and flag not in (1, 2, 3):
            continue
        if name in used:
            used[name] += 1
            name = "%s_%d" % (name, used[name])
        else:
            used[name] = 1
        columns.append((index, name))
    if not columns:
        return None
    data = []
    for raw in rows[3:]:
        obj = {}
        for index, name in columns:
            if index >= len(raw):
                continue
            value = cell_value(raw[index])
            if value is None:
                continue
            obj[name] = value
        if obj:
            data.append(obj)
    return data


def pick_primary(stem: str, sheets: list[dict]) -> dict:
    def rank(sheet: dict) -> tuple:
        name = str(sheet["sheet"]).casefold().replace(" ", "")
        key = stem.casefold()
        if name == key or name == key + ".csv":
            band = 3
        elif key in name:
            band = 2
        else:
            band = 1
        return (band, len(sheet["rows"]))

    return max(sheets, key=rank)


def build_system_gate(rows: list[dict]) -> list[dict]:
    gates = []
    for row in rows:
        key = row.get("key")
        if not isinstance(key, str) or not key:
            continue
        if "open_day" not in key and not key.endswith("_days"):
            continue
        day = as_int(row.get("value"))
        if day is None:
            continue
        item = {"key": key, "open_server_day": day}
        remark = row.get("remark")
        if isinstance(remark, str) and remark:
            item["note"] = remark
        gates.append(item)
    gates.sort(key=lambda item: item["key"])
    return gates


def project_opencondition(rows: list[dict], gates: list[dict]) -> list[dict]:
    days = {}
    for gate in gates:
        key = gate["key"]
        if key.endswith("_open_day"):
            days[key[: -len("_open_day")]] = gate["open_server_day"]
    projected = []
    for row in rows:
        atom = row.get("atom")
        if atom is None or atom == "":
            continue
        atom_text = str(atom)
        level = as_int(row.get("condition_value"))
        kind = as_int(row.get("condition_type"))
        note = row.get("note")
        name = note if isinstance(note, str) and note else atom_text
        item = dict(row)
        lang = item.get("name")
        item["id"] = atom_text
        item["name"] = name
        if lang is not None and lang != name:
            item["name_lang"] = lang
        item["condition_type"] = 0 if kind is None else kind
        item["level"] = 0 if level is None else level
        item["open_server_day"] = days.get(atom_text, 0)
        projected.append(item)
    alias(projected, "main_pub_task", "main_cafe")
    alias(projected, "main_mail", "mail")
    ensure_nav_bag(projected)
    return projected


def ensure_nav_bag(rows: list[dict]) -> None:
    """City nav uses id bag. The sheet has backpack_* atoms, not bag."""
    if any(row.get("id") == "bag" for row in rows):
        return
    rows.append(
        {
            "id": "bag",
            "name": "背包",
            "condition_type": 1,
            "level": 1,
            "open_server_day": 0,
            "note": "客户端导航 id。opencondition 没有 bag 原子，等级 1 与原来的入口门槛一致。",
        }
    )


def alias(rows: list[dict], source_id: str, alias_id: str) -> None:
    if any(row.get("id") == alias_id for row in rows):
        return
    for row in rows:
        if row.get("id") == source_id:
            copy = dict(row)
            copy["id"] = alias_id
            copy["alias_of"] = source_id
            rows.append(copy)
            return


def write_table(out: Path, name: str, note: str, rows: list[dict]) -> None:
    path = out / ("%s.json" % name)
    write_json(path, {"table": name, "note": note, "rows": rows})
    ensure_meta(Path(str(path) + ".meta"))


def write_json(path: Path, payload: dict) -> None:
    text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    path.write_text(text + "\n", encoding="utf-8")


def ensure_meta(path: Path) -> None:
    if path.exists():
        return
    path.write_text(META_TEMPLATE.format(uuid=str(uuid.uuid4())), encoding="utf-8")


def sheet_slug(name: str, index: int) -> str:
    chars = []
    for ch in name:
        if ("a" <= ch <= "z") or ("A" <= ch <= "Z") or ("0" <= ch <= "9") or ch in "-_":
            chars.append(ch)
        else:
            chars.append("_")
    text = "".join(chars).strip("_")
    while "__" in text:
        text = text.replace("__", "_")
    if not text:
        text = "sheet%d" % index
    return text[:48]


def is_type_word(value) -> bool:
    return isinstance(value, str) and value.strip().lower() in TYPE_WORDS


def is_field(value) -> bool:
    if not isinstance(value, str):
        return False
    return FIELD_RE.match(value.strip()) is not None


def flag_of(value):
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and value in (0, 1, 2, 3):
        return value
    if isinstance(value, str) and value.strip() in ("0", "1", "2", "3"):
        return int(value.strip())
    return "other"


def as_int(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and INT_RE.match(value.strip()):
        return int(value.strip())
    return None


def cell_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        if value.is_integer():
            return int(value)
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return None
    return text


if __name__ == "__main__":
    sys.exit(main())
